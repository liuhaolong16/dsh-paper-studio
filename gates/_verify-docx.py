#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
_verify-docx.py —— **成品件核验器**：重新打开生成的 docx，数原生对象并扫残留标记。

为什么要单独设一台（方法论，别删）：
  **源文件全绿 ≠ 成品全绿**——两者的缺陷空间不同：源文件里错的是**内容**，
  成品里错的是**格式承载**（标记没被渲染、表格退化成文本、图片没嵌进去）。
  所以 md 侧的门全过，也必须再打开成品数一遍。

判据（可数对象，不靠肉眼看）：
  · **通用不变量**：正文/表格单元格里不得残留 Markdown 标记（`**` / 反引号 / 行首 `|`）
  · **可选期望**：--expect-tables N / --expect-images N（给了才断言；没给只报数）

用法：
  python _verify-docx.py <输出.docx> [--expect-tables N] [--expect-images N] [--report 路径]

退出码：0 通过 / 1 违反不变量（有残留标记） / 2 与期望数不符 / 10 参数错
"""
import json
import os
import re
import sys

# 本机控制台是 GBK：直接 print 中文会出现乱码（下游按关键字抓输出就会抓不到）。
# 与其它脚本统一：把 stdout/stderr 显式切到 utf-8（errors=replace，不因编码崩掉）。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:                                            # noqa: BLE001
    pass

from docx import Document

argv = sys.argv[1:]
exp_t = exp_i = None
report = None
for key, setter in (("--expect-tables", "t"), ("--expect-images", "i")):
    if key in argv:
        i = argv.index(key)
        val = int(argv[i + 1])
        del argv[i:i + 2]
        if setter == "t":
            exp_t = val
        else:
            exp_i = val
if "--report" in argv:
    i = argv.index("--report")
    report = argv[i + 1]
    del argv[i:i + 2]
JSON_OUT = "--json" in argv
argv = [a for a in argv if a != "--json"]

pos = [a for a in argv if not a.startswith("--")]
if not pos:
    print("用法：python _verify-docx.py <docx> [--expect-tables N] [--expect-images N] [--report 路径]")
    sys.exit(10)

path = pos[0]
if not os.path.exists(path):
    print("找不到文件：%s" % path)
    sys.exit(10)

d = Document(path)
paras = [p for p in d.paragraphs if p.text.strip()]
heads = [p for p in d.paragraphs if p.style.name == "Title" or p.style.name.startswith("Heading")]
imgs = len(d.inline_shapes)
tabs = len(d.tables)

BAD = re.compile(r"\*\*|`|^\s*\|")
leak_p = [p.text[:70] for p in d.paragraphs if BAD.search(p.text)]
leak_c = []
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if BAD.search(c.text):
                leak_c.append(c.text[:50])

words = sum(len(p.text.split()) for p in paras)

print("文件            %s  (%.1f KB)" % (path, os.path.getsize(path) / 1024.0))
print("非空段落        %d" % len(paras))
print("标题数          %d" % len(heads))
print("原生表格        %d%s" % (tabs, "" if exp_t is None else "（期望 %d）" % exp_t))
print("内嵌图片        %d%s" % (imgs, "" if exp_i is None else "（期望 %d）" % exp_i))
print("正文词数        %d" % words)
print("残留 Markdown  段落 %d 处 / 表格单元格 %d 处" % (len(leak_p), len(leak_c)))
for x in leak_p[:5]:
    print("   ! " + x)
for x in leak_c[:5]:
    print("   ! " + x)

code = 0
if leak_p or leak_c:
    print("判定            失败：成品里残留 Markdown 标记（说明转换器漏渲染）")
    code = 1
elif (exp_t is not None and tabs != exp_t) or (exp_i is not None and imgs != exp_i):
    print("判定            失败：原生对象数与期望不符（表格/图片可能在转换中丢了）")
    code = 2
else:
    print("判定            通过")

if JSON_OUT:
    # 机器可读的一行（前缀便于在下游把人的输出与它分开）：供 `studio crosscheck` 与另一个
    # **独立实现**（插件里 ZIP+XML 层自己数）逐字段对拍 —— 两个实现判同一份文件必须同值。
    print("##JSON## " + json.dumps({
        "tables": tabs, "inline_images": imgs,
        "leak_paragraphs": len(leak_p), "leak_cells": len(leak_c),
        "paragraphs": len(paras), "headings": len(heads), "words": words,
        "bytes": os.path.getsize(path), "exit": code,
    }, ensure_ascii=False))

if report:
    with open(report, "w", encoding="utf-8") as fh:
        fh.write("tables=%d images=%d leak_paras=%d leak_cells=%d words=%d code=%d\n"
                 % (tabs, imgs, len(leak_p), len(leak_c), words, code))
sys.exit(code)
