#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""隐私与保密门（paper-studio 第 11 道 · V1.7）

为什么有这道门：此前的门管"声明齐不齐 / 工作注记有没有越界 / 引用核不核得到 / 方法要件全不全"，
但**没有一道管"这段内容该不该出现在投稿稿里"**。审稿人/编辑视角下这是硬项：
  · 审稿意见痕迹混进正文（投稿稿不是 response letter）
  · 未脱敏的个人信息 / 患者标识
  · 内部路径与内部材料名（本机路径、共享盘、内网）
  · 他人未发表内容未获授权
出处（本地精读的一篇 OA 论文，Lee et al. 2025, JEEHP 22:4，第五节「Key challenges」）：
  隐私与保密（privacy and confidentiality）、**过早数据披露**（premature data disclosure）
  被列为 LLM 进入审稿流程的两大风险之一；本门把"作者侧"该守的几条变成可机检判据。

★ 分级（本门是**第一个带两级严重度**的门）：HARD ⇒ 退出码 1；只有 SOFT ⇒ 退出码 0 但打印提示。
★ 不重叠纪律：伦理/同意**声明在不在**由 declaration-set 管不到（它只管四段声明），所以本门只管
  "出现受试者类字眼却没有任何伦理/同意表述"这一条，且列为 SOFT（各刊处理不同）。
退出码：0 = 无 HARD（可有 SOFT）；1 = 有 HARD；2 = 前提层失败

用法：python _gate-confidentiality.py --draft <稿.md> [--docx <成品.docx>] [--json] [--selftest]
"""
import argparse
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:                                            # noqa: BLE001
    pass

# ── HARD：审稿/编辑通信痕迹（作者回应语气，投稿正文里不该有）──
REVIEWER_TRACES = re.compile(
    r"we thank the (anonymous )?reviewer|as (suggested|requested) by (the )?reviewer|"
    r"reviewer\s*#?\s*[123]\b|the (editor|referee) (suggested|asked|requested)|"
    r"response to (the )?review|decision letter|referee report|"
    r"审稿人(建议|要求|意见)|编辑(建议|要求)|修改意见(已|如下)|答复审稿", re.I)

# ── HARD：内部路径 / 内部材料名 ──
INTERNAL_PATHS = re.compile(
    r"[A-Za-z]:\\Users\\|\\\\[A-Za-z0-9._-]+\\[A-Za-z0-9._$-]+|"
    r"内网|公司内部|内部资料|内部共享|internal use only|confidential internal|"
    r"未公开数据|内部数据|私有仓库", re.I)

# ── HARD：敏感个人标识（身份证 / 手机号 / 病历号 / 学号）──
SENSITIVE_IDS = re.compile(
    r"\b1[3-9]\d{9}\b|\b\d{17}[\dXx]\b|"
    r"病历号|病案号|患者(编号|标识|ID)|受试者(编号|标识)|住院号|学号\s*[:：]?\s*\d{6,}", re.I)

# ── SOFT：联系方式（投稿稿里正常存在，只提示"确认是否应放正文"）──
CONTACT = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]{2,}", re.I)

# ── SOFT：他人未发表内容 ──
UNPUBLISHED = re.compile(
    r"unpublished (data|manuscript|results)|in preparation|personal communication|"
    r"未发表(数据|稿件)|待发表|in press", re.I)

# ── SOFT：出现受试者类字眼但全文无伦理/同意表述 ──
HUMAN_SUBJECTS = re.compile(
    r"human (participants|subjects|volunteers)|patients?|clinical trial|informed consent|"
    r"受试者|人类受试者|患者|临床试验|知情同意", re.I)
ETHICS_WORDS = re.compile(
    r"ethic[s]? (approval|committee|review board)|IRB\b|institutional review board|"
    r"informed consent| Helsinki|伦理(委员会|审查|批准)|知情同意", re.I)


# ── SOFT：AI 使用声明是否说明了「未发表数据怎么处理」（数据是否外发给第三方服务）──
#    出处：本地精读的 Lee et al. 2025 JEEHP §Privacy and confidentiality —— 该文把
#    「pre-publication data protection / model data retention / healthcare compliance」列为三大关切，
#    即：作者必须说清**未发表数据有没有进第三方 LLM 服务**（医疗刊多按硬项处理）。
#    不重叠纪律：declaration-set 只判「AI 声明在不在」，本门補的是「声明里有没有交代数据处理」。
AI_USE = re.compile(
    r"(生成式\s*AI|generative\s+AI|AI[- ]assisted|AI\s+use|large\s+language\s+model|LLM|ChatGPT|大语言模型|人工智能辅助)[^\n]{0,120}?"
    r"(声明|declaration|statement|润色|polish|起草|draft|辅助|assistance|使用)", re.I)
DATA_HANDLING = re.compile(
    r"(未上传|不上传|未外发|不外发|本地(处理|运行|部署|存储)|离线|已脱敏|脱敏处理|不涉及未发表数据|"
    r"no data (were|was) (uploaded|shared|transferred)|"
    r"(?:was|were|is|are)?\s*not\s+(?:been\s+)?(?:uploaded|shared|sent|transferred|disclosed)|"
    r"local(?:ly)? (?:processed|run|deployed|stored)|"
    r"(?:processed|analysed|analyzed|run|stored|handled) locally|"
    r"de-?identif|anonymi[sz]ed|on-?premise)", re.I)


def scan(text, docx_text=""):
    body = text + ("\n" + docx_text if docx_text else "")
    lines = body.splitlines()
    hard, soft = [], []

    def collect(rx, sev, rule, why):
        for i, ln in enumerate(lines, 1):
            for m in rx.finditer(ln):
                item = {"line": i, "rule": rule, "severity": sev,
                        "text": ln.strip()[:160], "hit": m.group(0)[:60], "why": why}
                (hard if sev == "HARD" else soft).append(item)

    collect(REVIEWER_TRACES, "HARD", "reviewer-trace",
            "投稿正文里出现了审稿/编辑通信痕迹 —— 这类内容属于 response letter，不在投稿稿里")
    collect(INTERNAL_PATHS, "HARD", "internal-path",
            "出现了内部路径/内部材料名 —— 公开稿件里引用内部位置会同时泄露环境与流程")
    collect(SENSITIVE_IDS, "HARD", "sensitive-id",
            "出现敏感个人标识（身份证/手机号/病历号/患者或受试者编号）—— 必须脱敏")
    collect(CONTACT, "SOFT", "contact",
            "出现邮箱等联系方式 —— 若是作者信息则正常，但请确认它不属于应放投稿系统的内容")
    collect(UNPUBLISHED, "SOFT", "unpublished-source",
            "引用了未发表内容 —— 请确认已获许可，并在正文写明来源与授权状态")
    if HUMAN_SUBJECTS.search(body) and not ETHICS_WORDS.search(body):
        soft.append({"line": None, "rule": "ethics-silent", "severity": "SOFT", "text": "", "hit": "",
                     "why": "正文提到受试者/患者/临床试验，但**全稿未见伦理批准或知情同意表述** —— "
                            "多数医学刊为硬项，请按目标刊要求补齐"})
    # AI 使用声明必须交代「未发表数据怎么处理」（外部证据：JEEHP 2025 §Privacy，三大关切之首）
    if AI_USE.search(body) and not DATA_HANDLING.search(body):
        soft.append({"line": None, "rule": "ai-data-handling", "severity": "SOFT", "text": "", "hit": "",
                     "why": "写了 AI 使用，但**没说未发表数据有没有外发给第三方模型服务** —— 建议补一句："
                            "「未发表数据未上传任何第三方服务 / 全部本地处理 / 已脱敏」；医学与部分出版社会按硬项要求"})
    return hard, soft


def check(draft, docx=None):
    if not draft or not os.path.exists(draft):
        return 2, {"error": "源稿不存在：%s" % draft}
    text = open(draft, encoding="utf-8", errors="replace").read()
    dx = ""
    if docx:
        if not os.path.exists(docx):
            return 2, {"error": "成品件不存在：%s" % docx}
        try:
            from docx import Document
            d = Document(docx)
            dx = "\n".join(p.text for p in d.paragraphs)
            dx += "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        except Exception as e:                               # noqa: BLE001
            return 2, {"error": "成品件读取失败（前提层）：%s" % e}
    hard, soft = scan(text, dx)
    return (1 if hard else 0), {"schema": "confidentiality.v1", "draft": draft, "docx": docx or "",
                                "hard": hard, "soft": soft,
                                "findings": hard, "notes": [s["why"] for s in soft],
                                "verdict": "FAIL" if hard else "PASS"}


CLEAN = """# Cross-temperature generalisation of battery SOH prediction
Abstract: We measured six cells and report error bars.
# 2. Method
We trained three backbones on a public dataset and released our code.
# 5. Limitations
Our study has limitations. Internal validity was checked by a held-out split.
# Declarations
Competing interests: none. Funding: none. Ethics: not applicable (no human participants).
"""


def selftest():
    import tempfile
    tmp = tempfile.mkdtemp(prefix="conf-")

    def w(name, body):
        p = os.path.join(tmp, name)
        open(p, "w", encoding="utf-8").write(body)
        return p

    rev = "# Paper\n# 3. Response\nWe thank the reviewer for the suggestion; we have revised.\n"
    path = "# Paper\n# 2. Method\nRaw data lives at C:\\Users\\ASUS\\data\\run3.csv\n"
    pid = "# Paper\n# 4. Data\nPatient 编号: 患者编号 00123 and 学号: 202112345\n"
    mail = CLEAN + "\nCorresponding author: someone@example.org\n"
    hum = "# Paper\n# 4. Results\nWe recruited 40 patients for the trial.\n"
    unp = CLEAN + "\nWe compare against unpublished data from a partner lab.\n"
    peer = CLEAN + "\nPeer review is a cornerstone of science, and reviewers vary in strictness.\n"
    iv = CLEAN + "\nInternal validity, external validity and statistical conclusion validity differ.\n"
    cases = [
        ("① 干净稿必须过（负控）", w("a.md", CLEAN), 0),
        ("② 审稿通信痕迹 ⇒ HARD 失败", w("b.md", rev), 1),
        ("③ 本机内部路径 ⇒ HARD 失败", w("c.md", path), 1),
        ("④ 敏感个人标识 ⇒ HARD 失败", w("d.md", pid), 1),
        ("⑤ 联系邮箱 ⇒ 只 SOFT（不失败）", w("e.md", mail), 0),
        ("⑥ 提到患者却无伦理表述 ⇒ SOFT（不失败）", w("f.md", hum), 0),
        ("⑦ 引用未发表内容 ⇒ SOFT（不失败）", w("g.md", unp), 0),
        ("⑧ ★反向控制：正常讨论同行评审 ⇒ 不得报 HARD", w("h.md", peer), 0),
        ("⑨ ★反向控制：internal validity 这类统计术语 ⇒ 不得当内部路径", w("i.md", iv), 0),
        ("⑩ 缺文件 ⇒ 前提层失败", None, 2),
        ("⑫ AI 使用声明未交代数据处理 ⇒ SOFT（不失败）", w("n.md", CLEAN + "\n# Declaration of generative AI\nDuring preparation the authors used a large language model for language polishing.\n"), 0),
        ("⑬ ★反向控制：写明本地处理/未外发 ⇒ 不得再提示数据处理", w("o.md", CLEAN + "\n# Declaration of generative AI\nWe used a large language model for language polishing only; all unpublished data were processed locally and were not uploaded to any third-party service.\n"), 0),
    ]
    print("════ _gate-confidentiality --selftest ════")
    bad = 0
    for name, p, want in cases:
        code, payload = check(p)
        ok = (code == want)
        if not ok:
            bad += 1
        print("  %s %s%s" % ("✔" if ok else "✘", name, "" if ok else "  ← 期望 %s 实得 %s（%s）"
                             % (want, code, payload.get("error") or payload.get("verdict"))))
    # 额外断言：⑤⑥⑦ 必须是 SOFT 而不是 HARD（分级不许糊）
    _, p5 = check(os.path.join(tmp, "e.md"))
    _, p7 = check(os.path.join(tmp, "g.md"))
    graded = (not p5["hard"] and any(s["rule"] == "contact" for s in p5["soft"])
              and not p7["hard"] and any(s["rule"] == "unpublished-source" for s in p7["soft"]))
    if not graded:
        bad += 1
    print("  %s ⑪ 分级：邮箱与未发表内容必须是 SOFT 而不是 HARD" % ("✔" if graded else "✘"))
    # 额外断言：AI 数据处理提示（该提示时必须提示；写明本地处理后必须不再提示）
    _, p12 = check(os.path.join(tmp, "n.md"))
    _, p13 = check(os.path.join(tmp, "o.md"))
    dh = (not p12["hard"] and any(s["rule"] == "ai-data-handling" for s in p12["soft"])
          and not any(s["rule"] == "ai-data-handling" for s in p13["soft"]))
    if not dh:
        bad += 1
    print("  %s ⑭ AI 数据处理提示：该提示时必须提示、写明本地处理后必须不再提示" % ("✔" if dh else "✘"))
    total = len(cases) + 2                                        # 用例 + 两条跨用例断言
    print("\n  %s" % ("✔ 夹具 %d/%d 判别正确（用例 %d + 跨用例断言 2）" % (total - bad, total, len(cases))
                      if not bad else "✘ 夹具失败 %d/%d —— 本门当前不可信" % (bad, total)))
    return 1 if bad else 0


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("--draft")
    ap.add_argument("--docx")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(selftest())
    if not a.draft:
        print("用法：--draft <稿.md> [--docx <成品.docx>] [--json] ｜ --selftest")
        sys.exit(2)
    code, payload = check(a.draft, a.docx)
    if a.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        sys.exit(code)
    if code == 2:
        print("前提层失败：%s" % payload.get("error"))
        sys.exit(2)
    print("════ 隐私与保密门（第 11 道 · HARD/SOFT 两级） ════")
    for f in payload["hard"]:
        print("  [HARD] L%s %s：%s" % (f["line"], f["rule"], f["hit"]))
        print("         %s" % f["text"])
    for s in payload["soft"]:
        loc = ("L%s" % s["line"]) if s["line"] else "全稿"
        print("  [SOFT] %s %s：%s" % (loc, s["rule"], s["why"]))
    if payload["hard"]:
        print("\n  判定：失败 —— %d 条硬项；投稿前必须处理（隐私/保密是审稿硬项）" % len(payload["hard"]))
    else:
        print("\n  判定：通过（无硬项%s）" % ("；%d 条提示待你确认" % len(payload["soft"]) if payload["soft"] else ""))
    sys.exit(code)


if __name__ == "__main__":
    main()
